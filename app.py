import json
import re
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

here = Path(__file__).parent
HEADER = ["Español", "English", "Gender", "Added"]

st.set_page_config(page_title="El o La", page_icon="🃏", layout="centered")
st.markdown(
    "<style>.block-container{padding-top:1rem} header{visibility:hidden}</style>",
    unsafe_allow_html=True,
)


def secret(key):
    try:
        return st.secrets[key]
    except Exception:
        return None


@st.cache_data
def excel_words():
    wb = load_workbook(here / "Genero_Gramatical.xlsx", read_only=True)
    words = []
    for gender in ("El", "La"):
        for es, en in wb[gender].iter_rows(min_row=2, values_only=True):
            if es:
                words.append({"es": es.strip(), "en": (en or "").strip(), "g": gender})
    return words


@st.cache_resource
def worksheet():
    creds, url = secret("gcp_service_account"), secret("sheet_url")
    if not creds or not url:
        return None
    import gspread
    return gspread.service_account_from_dict(dict(creds)).open_by_url(url).sheet1


@st.cache_data(ttl=60)
def sheet_words():
    ws = worksheet()
    if ws is None:
        return []
    words = []
    for row_no, r in enumerate(ws.get_all_values()[1:], start=2):
        es, en, g = [(c or "").strip() for c in (r + ["", "", ""])[:3]]
        g = g.capitalize()
        if es and g in ("El", "La"):
            words.append({"es": es, "en": en, "g": g, "row": row_no})
    return words


@st.cache_data
def build_page(words_json):
    html = (here / "flashcards_template.html").read_text(encoding="utf-8")
    html = html.replace("/*DATA*/", words_json)
    # The in-page add panel needs claude.ai storage; this app adds words through the form below instead.
    return html.replace("</style>", "#add{display:none!important}</style>", 1)


def parse_notes(text, known):
    multi = sorted((w for w in known if " " in w), key=len, reverse=True)
    out, g = [], None
    for line in text.splitlines():
        for raw in line.split("/"):
            seg = " ".join(raw.split())
            low = seg.lower()
            if not seg:
                continue
            if re.fullmatch(r"(el|la)[:.]?", low):
                g = "El" if low.startswith("el") else "La"
                continue
            if low in ("un", "una"):
                continue
            seg = re.sub(r"^(un|una)\s+", "", seg, flags=re.I)
            row_g = g
            art = re.match(r"^(el|la)\s+(.+)$", seg, re.I)
            if art:
                row_g, seg = ("El" if art[1].lower() == "el" else "La"), art[2]
            sep = re.match(r"^(.+?)\s*(?:=|:|\s[-–—]\s)\s*(.*)$", seg)
            k = next((p for p in multi if seg.lower() == p.lower() or seg.lower().startswith(p.lower() + " ")), None)
            de = re.match(r"^(\S+ de \S+)(?:\s+(.*))?$", seg, re.I)
            if sep:
                es, en = sep[1], sep[2]
            elif k:
                es, en = seg[: len(k)], seg[len(k):]
            elif de:
                es, en = de[1], de[2] or ""
            else:
                es, _, en = seg.partition(" ")
            out.append({"Gender": row_g, "Español": es.strip(), "English": en.strip()})
    return out


excel = excel_words()
try:
    added = sheet_words()
    sheet_error = None
except Exception as e:
    added, sheet_error = [], e
base = {w["es"].lower() for w in excel}
all_words = excel + [{k: w[k] for k in ("es", "en", "g")} for w in added if w["es"].lower() not in base]

st.iframe(build_page(json.dumps(all_words, ensure_ascii=False)), height=1200)


def add_words_panel():
    ws = worksheet()
    if ws is None:
        st.caption("Adding words isn't set up yet. It needs a Google Sheet connected in the app's secrets.")
        return
    if sheet_error:
        st.error("The Google Sheet couldn't be read, so adding words is paused. Check that the sheet is shared with the service account.")
        return

    pw = secret("add_password")
    if pw and not st.session_state.get("unlocked"):
        entered = st.text_input("Password to add words", type="password", key="pw")
        if entered == pw:
            st.session_state.unlocked = True
            st.rerun()
        elif entered:
            st.error("That password isn't right.")
        return

    st.write("Paste words the way you write your notes. Start with **La** or **El**, then separate words with slashes. The English is optional.")
    text = st.text_area("Words to add", key="paste", placeholder="La/una mesa table/calle street/ventana\nEl/un libro book/perro dog/problema", label_visibility="collapsed")
    if st.button("Preview"):
        st.session_state.preview = parse_notes(text, [w["es"] for w in all_words])
        if not st.session_state.preview:
            st.info("Nothing to preview yet. Paste some words first.")

    rows = st.session_state.get("preview") or []
    if rows:
        known = {w["es"].lower(): w for w in all_words}
        df = pd.DataFrame(rows, columns=["Gender", "Español", "English"])
        edited = st.data_editor(
            df,
            key="editor",
            hide_index=True,
            width="stretch",
            num_rows="dynamic",
            column_config={"Gender": st.column_config.SelectboxColumn("el / la", options=["El", "La"])},
        )
        ready, dups, unset, seen = [], [], [], set()
        for r in edited.fillna("").to_dict("records"):
            es = str(r["Español"]).strip()
            if not es:
                continue
            if es.lower() in known or es.lower() in seen:
                dups.append(es)
            elif r["Gender"] not in ("El", "La"):
                unset.append(es)
            else:
                ready.append([es, str(r["English"]).strip(), r["Gender"]])
            seen.add(es.lower())
        if dups:
            st.caption("Already in your deck, will be skipped: " + ", ".join(dups))
        if unset:
            st.warning("Pick el or la for: " + ", ".join(unset))
        c1, c2 = st.columns(2)
        if c1.button(f"Add {len(ready)} word{'s' if len(ready) != 1 else ''}", type="primary", disabled=not ready):
            now = datetime.now(timezone.utc).isoformat(timespec="seconds")
            try:
                if not ws.row_values(1):
                    ws.append_row(HEADER)
                ws.append_rows([r + [now] for r in ready], value_input_option="RAW")
            except Exception:
                st.error("The words couldn't be saved to the Google Sheet. Try again.")
            else:
                st.session_state.preview = []
                st.session_state.pop("editor", None)
                st.session_state.flash = f"Added {len(ready)} word{'s' if len(ready) != 1 else ''}. They'll show up in your next rounds."
                sheet_words.clear()
                st.rerun()
        if c2.button("Discard"):
            st.session_state.preview = []
            st.session_state.pop("editor", None)
            st.rerun()

    if st.session_state.get("flash"):
        st.success(st.session_state.pop("flash"))

    if added:
        with st.expander(f"Words you've added ({len(added)})"):
            df = pd.DataFrame([{"Remove": False, "el / la": w["g"], "Español": w["es"], "English": w["en"]} for w in added])
            out = st.data_editor(df, key="mine", hide_index=True, width="stretch",
                                 disabled=["el / la", "Español", "English"])
            picked = [added[i] for i in out.index[out["Remove"]]]
            if st.button(f"Remove {len(picked)} selected", disabled=not picked):
                try:
                    for w in sorted(picked, key=lambda w: w["row"], reverse=True):
                        ws.delete_rows(w["row"])
                except Exception:
                    st.error("Some words couldn't be removed. Reload and try again.")
                st.session_state.pop("mine", None)
                sheet_words.clear()
                st.rerun()


st.subheader("Add words")
add_words_panel()
